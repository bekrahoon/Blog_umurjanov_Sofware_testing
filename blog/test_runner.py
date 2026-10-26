"""
Свой test runner для Django: прогоняет ВСЕ тесты проекта (`python manage.py test`)
и после прогона записывает отчёт в Word-файл Test_Report.docx.

Подключается в settings.py:
    TEST_RUNNER = 'blog.test_runner.WordReportRunner'

В отчёте:
  1) сводка (сколько тестов, сколько прошло / упало);
  2) таблица: №, класс, тест, что проверяет (docstring), результат;
  3) код каждого теста.
"""
import datetime
import inspect
import textwrap
import unittest
from pathlib import Path

from django.conf import settings
from django.test.runner import DiscoverRunner
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

STATUS_TEXT = {
    'success': 'Успешно',
    'failure': 'Провал',
    'error': 'Ошибка',
    'skip': 'Пропущен',
}
STATUS_COLOR = {
    'success': RGBColor(0x1E, 0x7B, 0x34),
    'failure': RGBColor(0xC6, 0x28, 0x28),
    'error': RGBColor(0xC6, 0x28, 0x28),
    'skip': RGBColor(0x80, 0x80, 0x80),
}
STATUS_FILL = {
    'success': 'E6F4EA',
    'failure': 'FDE7E7',
    'error': 'FDE7E7',
    'skip': 'F1F1F1',
}


class WordReportResult(unittest.TextTestResult):
    """Обычный результат unittest, который дополнительно запоминает каждый тест."""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.records = []

    def _record(self, test, status, err=None, reason=''):
        method_name = getattr(test, '_testMethodName', str(test))
        detail = reason
        if err is not None:
            # последняя строка traceback, например "AssertionError: 404 != 200"
            detail = self._exc_info_to_string(err, test).strip().splitlines()[-1]

        source = ''
        method = getattr(type(test), method_name, None)
        if method is not None:
            try:
                source = textwrap.dedent(inspect.getsource(method))
            except (OSError, TypeError):
                pass

        self.records.append({
            'class': type(test).__name__,
            'name': method_name,
            'doc': (test.shortDescription() or '') if hasattr(test, 'shortDescription') else '',
            'status': status,
            'detail': detail,
            'source': source,
        })

    def addSuccess(self, test):
        super().addSuccess(test)
        self._record(test, 'success')

    def addFailure(self, test, err):
        super().addFailure(test, err)
        self._record(test, 'failure', err)

    def addError(self, test, err):
        super().addError(test, err)
        self._record(test, 'error', err)

    def addSkip(self, test, reason):
        super().addSkip(test, reason)
        self._record(test, 'skip', reason=reason)

    def addExpectedFailure(self, test, err):
        super().addExpectedFailure(test, err)
        self._record(test, 'success', reason='ожидаемый провал')

    def addUnexpectedSuccess(self, test):
        super().addUnexpectedSuccess(test)
        self._record(test, 'failure', reason='неожиданно прошёл')


class WordReportRunner(DiscoverRunner):
    """Стандартный runner Django + запись отчёта в Word после прогона."""

    def get_resultclass(self):
        return WordReportResult

    def run_suite(self, suite, **kwargs):
        result = super().run_suite(suite, **kwargs)
        path = Path(getattr(settings, 'TEST_REPORT_PATH', settings.BASE_DIR / 'Test_Report.docx'))
        build_report(result.records, path)
        print(f'\nОтчёт о тестировании сохранён: {path}')
        return result


# ---------- построение Word-документа ----------

def _shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_fill)
    tc_pr.append(shd)


def _cell_text(cell, text, bold=False, color=None, size=10):
    cell.text = ''
    run = cell.paragraphs[0].add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def build_report(records, path):
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Arial'
    normal.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('Отчёт об автоматическом тестировании')
    run.bold = True
    run.font.size = Pt(18)

    doc.add_paragraph('Проект: Django Blog (blog)')
    doc.add_paragraph('Студент: Умуржанов Аба-Бекрахун')
    doc.add_paragraph('Дата прогона: ' + datetime.datetime.now().strftime('%d.%m.%Y %H:%M'))
    doc.add_paragraph('Отчёт сформирован автоматически командой: python manage.py test')

    total = len(records)
    passed = sum(r['status'] == 'success' for r in records)
    failed = sum(r['status'] in ('failure', 'error') for r in records)
    skipped = sum(r['status'] == 'skip' for r in records)

    doc.add_heading('1. Сводка', level=1)
    summary = doc.add_table(rows=4, cols=2)
    summary.style = 'Table Grid'
    for row, (label, value) in zip(summary.rows, [
        ('Всего тестов', total),
        ('Успешно', passed),
        ('Провал / ошибка', failed),
        ('Пропущено', skipped),
    ]):
        _cell_text(row.cells[0], label, bold=True)
        _cell_text(row.cells[1], str(value))

    doc.add_heading('2. Результаты тестов', level=1)
    table = doc.add_table(rows=1, cols=5)
    table.style = 'Table Grid'
    for cell, text in zip(table.rows[0].cells, ['№', 'Класс', 'Тест', 'Что проверяет', 'Результат']):
        _cell_text(cell, text, bold=True)
        _shade(cell, 'D9E2F3')

    for i, rec in enumerate(records, start=1):
        cells = table.add_row().cells
        _cell_text(cells[0], str(i))
        _cell_text(cells[1], rec['class'])
        _cell_text(cells[2], rec['name'], size=9)
        _cell_text(cells[3], rec['doc'] or '—')
        status = STATUS_TEXT[rec['status']]
        if rec['detail']:
            status += '\n' + rec['detail']
        _cell_text(cells[4], status, bold=True, color=STATUS_COLOR[rec['status']])
        _shade(cells[4], STATUS_FILL[rec['status']])

    table.autofit = False
    widths = [Cm(1.0), Cm(3.2), Cm(4.6), Cm(5.7), Cm(2.5)]
    for column, width in zip(table.columns, widths):
        column.width = width
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = width

    doc.add_heading('3. Код тестов', level=1)
    for i, rec in enumerate(records, start=1):
        doc.add_heading(f"{i}. {rec['class']}.{rec['name']} — {STATUS_TEXT[rec['status']]}", level=3)
        code = doc.add_paragraph()
        code_run = code.add_run(rec['source'].rstrip() or '(код не найден)')
        code_run.font.name = 'Consolas'
        code_run.font.size = Pt(9)

    doc.save(path)
